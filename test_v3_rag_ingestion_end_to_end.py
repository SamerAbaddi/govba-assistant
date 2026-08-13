"""End-to-end offline regression tests for GovBA-GAR ingestion."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pymupdf
from docx import Document

from govba.rag import (
    AuthoritativeSource,
    ChunkingConfig,
    DocumentFormat,
    DocumentType,
    IngestionPipeline,
    IngestionRequest,
    SourceLanguage,
    SourceStatus,
)


def make_source(
    document_id: str,
    title: str,
    language: SourceLanguage = SourceLanguage.ENGLISH,
):
    return AuthoritativeSource(
        document_id=document_id,
        title=title,
        issuing_authority=(
            "Synthetic GovBA Test Authority"
        ),
        document_type=DocumentType.PROCEDURE,
        language=language,
        status=SourceStatus.CURRENT,
    )


def create_pdf(
    path: Path,
):
    document = pymupdf.open()

    try:
        page_one = document.new_page()

        page_one.insert_text(
            (72, 72),
            (
                "Annual leave requests require "
                "manager approval."
            ),
        )

        page_two = document.new_page()

        page_two.insert_text(
            (72, 72),
            (
                "Requests should be submitted "
                "five working days in advance."
            ),
        )

        document.save(path)

    finally:
        document.close()


def create_docx(
    path: Path,
):
    document = Document()

    document.add_heading(
        "Employee Services",
        level=1,
    )

    document.add_heading(
        "Annual Leave",
        level=2,
    )

    document.add_paragraph(
        "  Leave\trequests require\u00a0approval.  "
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(
        0,
        0,
    ).text = "Requirement"

    table.cell(
        0,
        1,
    ).text = "Deadline"

    table.cell(
        1,
        0,
    ).text = "Annual leave"

    table.cell(
        1,
        1,
    ).text = "Five working days"

    document.save(path)


class TestIngestionEndToEnd(
    unittest.TestCase
):
    def setUp(self):
        self.pipeline = IngestionPipeline(
            chunking_config=ChunkingConfig(
                target_characters=500,
                overlap_characters=50,
            )
        )

    def test_real_pdf_end_to_end(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "policy.pdf"

            create_pdf(path)

            request = IngestionRequest(
                source=make_source(
                    "DOC-E2E-PDF",
                    "Synthetic PDF Policy",
                ),
                file_path=path,
                document_format=DocumentFormat.PDF,
            )

            result = self.pipeline.ingest(
                request
            )

            self.assertGreaterEqual(
                result.block_count,
                2,
            )

            self.assertGreaterEqual(
                result.chunk_count,
                2,
            )

    def test_real_docx_end_to_end(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "policy.docx"

            create_docx(path)

            request = IngestionRequest(
                source=make_source(
                    "DOC-E2E-DOCX",
                    "Synthetic DOCX Policy",
                ),
                file_path=path,
                document_format=DocumentFormat.DOCX,
            )

            result = self.pipeline.ingest(
                request
            )

            self.assertGreaterEqual(
                result.block_count,
                4,
            )

            self.assertGreaterEqual(
                result.chunk_count,
                4,
            )

    def test_pdf_page_provenance_survives_pipeline(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "pages.pdf"

            create_pdf(path)

            result = self.pipeline.ingest(
                IngestionRequest(
                    source=make_source(
                        "DOC-PAGES",
                        "Page Provenance Test",
                    ),
                    file_path=path,
                    document_format=DocumentFormat.PDF,
                )
            )

            self.assertEqual(
                result.page_numbers,
                (
                    1,
                    2,
                ),
            )

            self.assertEqual(
                {
                    chunk.page_start
                    for chunk in result.chunks
                },
                {
                    1,
                    2,
                },
            )

    def test_docx_section_provenance_survives_pipeline(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sections.docx"

            create_docx(path)

            result = self.pipeline.ingest(
                IngestionRequest(
                    source=make_source(
                        "DOC-SECTIONS",
                        "Section Provenance Test",
                    ),
                    file_path=path,
                    document_format=DocumentFormat.DOCX,
                )
            )

            annual_leave_chunks = [
                chunk
                for chunk in result.chunks
                if chunk.section_title
                == "Annual Leave"
            ]

            self.assertTrue(
                annual_leave_chunks
            )

            for chunk in annual_leave_chunks:
                self.assertEqual(
                    chunk.section_path,
                    (
                        "Employee Services",
                        "Annual Leave",
                    ),
                )

    def test_normalization_survives_full_pipeline(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "normalized.docx"

            create_docx(path)

            result = self.pipeline.ingest(
                IngestionRequest(
                    source=make_source(
                        "DOC-NORMALIZED",
                        "Normalization Test",
                    ),
                    file_path=path,
                    document_format=DocumentFormat.DOCX,
                )
            )

            texts = tuple(
                chunk.text
                for chunk in result.chunks
            )

            self.assertIn(
                "Leave requests require approval.",
                texts,
            )

    def test_docx_table_survives_full_pipeline(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "table.docx"

            create_docx(path)

            result = self.pipeline.ingest(
                IngestionRequest(
                    source=make_source(
                        "DOC-TABLE",
                        "Table Test",
                    ),
                    file_path=path,
                    document_format=DocumentFormat.DOCX,
                )
            )

            combined = "\n".join(
                chunk.text
                for chunk in result.chunks
            )

            self.assertIn(
                "Requirement | Deadline",
                combined,
            )

            self.assertIn(
                "Annual leave | Five working days",
                combined,
            )

    def test_pdf_repeated_ingestion_is_deterministic(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "stable.pdf"

            create_pdf(path)

            request = IngestionRequest(
                source=make_source(
                    "DOC-STABLE-PDF",
                    "Stable PDF Test",
                ),
                file_path=path,
                document_format=DocumentFormat.PDF,
            )

            first = self.pipeline.ingest(
                request
            )

            second = self.pipeline.ingest(
                request
            )

            self.assertEqual(
                tuple(
                    chunk.chunk_id
                    for chunk in first.chunks
                ),
                tuple(
                    chunk.chunk_id
                    for chunk in second.chunks
                ),
            )

            self.assertEqual(
                tuple(
                    chunk.content_hash
                    for chunk in first.chunks
                ),
                tuple(
                    chunk.content_hash
                    for chunk in second.chunks
                ),
            )

    def test_docx_repeated_ingestion_is_deterministic(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "stable.docx"

            create_docx(path)

            request = IngestionRequest(
                source=make_source(
                    "DOC-STABLE-DOCX",
                    "Stable DOCX Test",
                ),
                file_path=path,
                document_format=DocumentFormat.DOCX,
            )

            first = self.pipeline.ingest(
                request
            )

            second = self.pipeline.ingest(
                request
            )

            self.assertEqual(
                tuple(
                    chunk.chunk_id
                    for chunk in first.chunks
                ),
                tuple(
                    chunk.chunk_id
                    for chunk in second.chunks
                ),
            )

    def test_source_identity_is_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "identity.pdf"

            create_pdf(path)

            source = make_source(
                "DOC-IDENTITY",
                "Identity Test",
            )

            result = self.pipeline.ingest(
                IngestionRequest(
                    source=source,
                    file_path=path,
                    document_format=DocumentFormat.PDF,
                )
            )

            self.assertEqual(
                result.extracted_document.source,
                source,
            )

            for chunk in result.chunks:
                self.assertEqual(
                    chunk.document_id,
                    source.document_id,
                )

    def test_one_pipeline_supports_both_formats(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)

            pdf_path = directory / "document.pdf"
            docx_path = directory / "document.docx"

            create_pdf(
                pdf_path
            )

            create_docx(
                docx_path
            )

            pdf_result = self.pipeline.ingest(
                IngestionRequest(
                    source=make_source(
                        "DOC-MIXED-PDF",
                        "Mixed PDF",
                    ),
                    file_path=pdf_path,
                    document_format=DocumentFormat.PDF,
                )
            )

            docx_result = self.pipeline.ingest(
                IngestionRequest(
                    source=make_source(
                        "DOC-MIXED-DOCX",
                        "Mixed DOCX",
                    ),
                    file_path=docx_path,
                    document_format=DocumentFormat.DOCX,
                )
            )

            self.assertGreater(
                pdf_result.chunk_count,
                0,
            )

            self.assertGreater(
                docx_result.chunk_count,
                0,
            )

            self.assertNotEqual(
                pdf_result.chunks[0].document_id,
                docx_result.chunks[0].document_id,
            )


if __name__ == "__main__":
    unittest.main()
