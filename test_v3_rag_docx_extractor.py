"""Offline tests for GovBA-GAR DOCX extraction."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from govba.rag import (
    AuthoritativeSource,
    DocumentExtractor,
    DocumentFormat,
    DocumentType,
    IngestionRequest,
    SourceLanguage,
    SourceStatus,
)
from govba.rag.docx_extractor import (
    DOCX_EXTRACTION_VERSION,
    DOCXExtractionError,
    PythonDocxExtractor,
)


def make_source():
    return AuthoritativeSource(
        document_id="DOC-DOCX-001",
        title="Synthetic DOCX Procedure",
        issuing_authority=(
            "Synthetic GovBA Test Authority"
        ),
        document_type=DocumentType.PROCEDURE,
        language=SourceLanguage.ENGLISH,
        status=SourceStatus.CURRENT,
    )


def make_request(
    path,
):
    return IngestionRequest(
        source=make_source(),
        file_path=Path(path),
        document_format=(
            DocumentFormat.DOCX
        ),
    )


class TestDOCXExtractor(unittest.TestCase):
    def test_version_is_stable(self):
        self.assertEqual(
            DOCX_EXTRACTION_VERSION,
            "govba-docx-extraction-v1",
        )

    def test_satisfies_extractor_protocol(self):
        self.assertIsInstance(
            PythonDocxExtractor(),
            DocumentExtractor,
        )

    def test_non_request_is_rejected(self):
        with self.assertRaises(TypeError):
            PythonDocxExtractor().extract(
                object()
            )

    def test_non_docx_request_is_rejected(self):
        request = IngestionRequest(
            source=make_source(),
            file_path=Path(
                "/tmp/document.pdf"
            ),
            document_format=(
                DocumentFormat.PDF
            ),
        )

        with self.assertRaises(ValueError):
            PythonDocxExtractor().extract(
                request
            )

    def test_missing_file_is_rejected(self):
        with self.assertRaises(
            FileNotFoundError
        ):
            PythonDocxExtractor().extract(
                make_request(
                    "/tmp/"
                    "definitely-missing.docx"
                )
            )

    def test_paragraph_is_extracted(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "paragraph.docx"
            )

            document = Document()
            document.add_paragraph(
                "Annual leave requires approval."
            )
            document.save(path)

            result = (
                PythonDocxExtractor()
                .extract(
                    make_request(path)
                )
            )

            self.assertEqual(
                result.block_count,
                1,
            )

            self.assertEqual(
                result.blocks[0].text,
                "Annual leave requires approval.",
            )

    def test_block_indexes_are_contiguous(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "indexes.docx"
            )

            document = Document()
            document.add_paragraph(
                "First paragraph."
            )
            document.add_paragraph(
                "Second paragraph."
            )
            document.save(path)

            result = (
                PythonDocxExtractor()
                .extract(
                    make_request(path)
                )
            )

            self.assertEqual(
                tuple(
                    block.block_index
                    for block in result.blocks
                ),
                tuple(
                    range(
                        result.block_count
                    )
                ),
            )

    def test_document_order_is_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "order.docx"
            )

            document = Document()

            document.add_paragraph(
                "Before table."
            )

            table = document.add_table(
                rows=1,
                cols=2,
            )

            table.cell(
                0,
                0,
            ).text = "Cell A"

            table.cell(
                0,
                1,
            ).text = "Cell B"

            document.add_paragraph(
                "After table."
            )

            document.save(path)

            result = (
                PythonDocxExtractor()
                .extract(
                    make_request(path)
                )
            )

            self.assertEqual(
                tuple(
                    block.text
                    for block in result.blocks
                ),
                (
                    "Before table.",
                    "Cell A | Cell B",
                    "After table.",
                ),
            )

    def test_heading_metadata_is_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "heading.docx"
            )

            document = Document()

            document.add_heading(
                "Eligibility",
                level=1,
            )

            document.add_paragraph(
                "Applicants must be eligible."
            )

            document.save(path)

            result = (
                PythonDocxExtractor()
                .extract(
                    make_request(path)
                )
            )

            body = result.blocks[1]

            self.assertEqual(
                body.section_title,
                "Eligibility",
            )

            self.assertEqual(
                body.section_path,
                (
                    "Eligibility",
                ),
            )

    def test_nested_heading_path_is_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "nested.docx"
            )

            document = Document()

            document.add_heading(
                "Services",
                level=1,
            )

            document.add_heading(
                "Eligibility",
                level=2,
            )

            document.add_paragraph(
                "Eligibility requirement."
            )

            document.save(path)

            result = (
                PythonDocxExtractor()
                .extract(
                    make_request(path)
                )
            )

            body = result.blocks[-1]

            self.assertEqual(
                body.section_path,
                (
                    "Services",
                    "Eligibility",
                ),
            )

            self.assertEqual(
                body.section_title,
                "Eligibility",
            )

    def test_table_is_extracted(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "table.docx"
            )

            document = Document()

            table = document.add_table(
                rows=2,
                cols=2,
            )

            table.cell(
                0,
                0,
            ).text = "Requirement"

            table.cell(
                0,
                1,
            ).text = "Deadline"

            table.cell(
                1,
                0,
            ).text = "Annual leave"

            table.cell(
                1,
                1,
            ).text = "Five days"

            document.save(path)

            result = (
                PythonDocxExtractor()
                .extract(
                    make_request(path)
                )
            )

            self.assertEqual(
                result.blocks[0].text,
                (
                    "Requirement | Deadline\n"
                    "Annual leave | Five days"
                ),
            )

    def test_source_metadata_is_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "metadata.docx"
            )

            document = Document()
            document.add_paragraph(
                "Policy content."
            )
            document.save(path)

            request = make_request(
                path
            )

            result = (
                PythonDocxExtractor()
                .extract(request)
            )

            self.assertEqual(
                result.source,
                request.source,
            )

            self.assertEqual(
                result.source_file_name,
                "metadata.docx",
            )

            self.assertEqual(
                result.document_format,
                DocumentFormat.DOCX,
            )

    def test_language_is_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "language.docx"
            )

            document = Document()
            document.add_paragraph(
                "Synthetic text."
            )
            document.save(path)

            result = (
                PythonDocxExtractor()
                .extract(
                    make_request(path)
                )
            )

            self.assertEqual(
                result.blocks[0].language,
                SourceLanguage.ENGLISH,
            )

    def test_empty_docx_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "empty.docx"
            )

            Document().save(path)

            with self.assertRaises(
                DOCXExtractionError
            ):
                PythonDocxExtractor().extract(
                    make_request(path)
                )

    def test_invalid_docx_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = (
                Path(directory)
                / "invalid.docx"
            )

            path.write_text(
                "This is not a DOCX package.",
                encoding="utf-8",
            )

            with self.assertRaises(
                DOCXExtractionError
            ):
                PythonDocxExtractor().extract(
                    make_request(path)
                )


if __name__ == "__main__":
    unittest.main()
